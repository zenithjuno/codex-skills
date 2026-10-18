from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
import json
import os
from unittest.mock import patch
from xml.etree import ElementTree as ET
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import thai_math_docx_builder as builder
import thai_math_docx_qa as qa
from thai_math_docx_preferences import load_spacing_preferences

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

class ParagraphSpacingTests(unittest.TestCase):
    def test_body_table_header_footer_styles_and_defaults(self):
        doc=builder.new_document()
        body=builder.add_paragraph(doc,parts=[{'type':'text','text':'ทดสอบ'}])
        self.assertEqual(body.paragraph_format.space_after.pt,0)
        table=doc.add_table(rows=1,cols=1)
        targets=[body,table.cell(0,0).paragraphs[0],doc.sections[0].header.paragraphs[0],doc.sections[0].footer.paragraphs[0]]
        for p in targets:
            p.text='คงข้อความ'
            p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.line_spacing=1.15
            spacing=p._p.get_or_add_pPr().find(qn('w:spacing'))
            spacing.set(qn('w:beforeLines'),'100')
            spacing.set(qn('w:afterAutospacing'),'1')
        doc.styles['Heading 1'].paragraph_format.line_spacing=1.15
        doc.styles['Normal'].paragraph_format.space_after=Pt(4)
        gap=builder.add_paragraph(doc,parts=[{'type':'text','text':'เว้นหก'}],space_after=6)
        gap.paragraph_format.space_before=Pt(6)
        with tempfile.TemporaryDirectory() as tmp:
            path=builder.save_docx(doc,Path(tmp)/'spacing.docx')
            with zipfile.ZipFile(path) as z:
                roots={name:ET.fromstring(z.read(name)) for name in z.namelist() if name.startswith('word/') and name.endswith('.xml')}
            self.assertEqual(qa._audit_paragraph_spacing(roots),[])
            for name in ('word/document.xml','word/styles.xml','word/header1.xml','word/footer1.xml'):
                spacings=roots[name].findall(f'.//{W}pPr/{W}spacing')
                self.assertTrue(spacings,name)
                for spacing in spacings:
                    self.assertIn(spacing.get(W+'before'),('0','120'))
                    self.assertIn(spacing.get(W+'after'),('0','120'))
                    self.assertEqual(spacing.get(W+'line'),'240')
                    self.assertEqual(spacing.get(W+'lineRule'),'auto')
            self.assertEqual(gap.paragraph_format.space_before.pt,6)
            self.assertEqual(gap.paragraph_format.space_after.pt,6)
            for p in targets:self.assertEqual(p.text,'คงข้อความ')
        before=doc.element.xml
        builder.normalize_document_spacing(doc)
        self.assertEqual(before,doc.element.xml)

    def test_gate_rejects_nonstandard_gaps_and_line_rules(self):
        for attrs in ({'after':'80'},{'before':'160'},{'line':'276','lineRule':'auto'},{'line':'240','lineRule':'exact'},{'beforeLines':'100'},{'afterAutospacing':'1'}):
            root=ET.Element(W+'document')
            p=ET.SubElement(root,W+'p');props=ET.SubElement(p,W+'pPr')
            ET.SubElement(props,W+'spacing',{W+k:v for k,v in attrs.items()})
            self.assertTrue(qa._audit_paragraph_spacing({'word/document.xml':root}),attrs)

    def test_project_override_drives_builder_normalizer_and_gate_together(self):
        with tempfile.TemporaryDirectory() as tmp:
            config=Path(tmp)/'spacing.json'
            config.write_text(json.dumps({'schema_version':1,'default_before_pt':0,'default_after_pt':8,'gap_pt':8,'line_spacing':1.25}))
            with patch.dict(os.environ,{'THAI_MATH_DOCX_SPACING_CONFIG':str(config)}):
                doc=builder.new_document()
                body=builder.add_paragraph(doc,parts=[{'type':'text','text':'กำหนดจาก config'}])
                heading=builder.add_heading(doc,'หัวข้อ')
                self.assertEqual(body.paragraph_format.space_after.pt,8)
                self.assertEqual(heading.paragraph_format.space_after.pt,8)
                self.assertEqual(body.paragraph_format.line_spacing,1.25)
                body.paragraph_format.space_after=Pt(6)
                path=builder.save_docx(doc,Path(tmp)/'configured.docx')
                self.assertEqual(body.paragraph_format.space_after.pt,8)
                with zipfile.ZipFile(path) as z:
                    roots={name:ET.fromstring(z.read(name)) for name in z.namelist() if name.startswith('word/') and name.endswith('.xml')}
                self.assertEqual(qa._audit_paragraph_spacing(roots),[])
                spacing=roots['word/document.xml'].find(f'.//{W}pPr/{W}spacing')
                spacing.set(W+'after','120')
                self.assertTrue(qa._audit_paragraph_spacing(roots))
            self.assertEqual(load_spacing_preferences()['line_spacing'],1.0)

    def test_invalid_config_does_not_silently_fall_back(self):
        base={'schema_version':1,'default_before_pt':0,'default_after_pt':0,'gap_pt':6,'line_spacing':1.0}
        for bad in ({'line_spacing':0},{'gap_pt':-1},{'default_after_pt':4},{'gap_pt':True},{'schema_version':2}):
            with tempfile.TemporaryDirectory() as tmp:
                config=Path(tmp)/'bad.json'
                config.write_text(json.dumps({**base,**bad}))
                with self.assertRaises(ValueError):load_spacing_preferences(config)

if __name__=='__main__':unittest.main()
