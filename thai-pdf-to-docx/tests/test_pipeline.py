"""End-to-end tests on a SYNTHETIC Thai PDF (built in tests/synth.py).  Run:  python3 -m pytest tests  (or: python3 -m unittest discover tests)

Skipped when no TH Sarabun New TTF is installed or the thai-font-normalize fixer is missing.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE))

import synth  # noqa: E402
from tdlib import find_fixer, find_soffice  # noqa: E402
import pipeline  # noqa: E402


@unittest.skipUnless(synth.fonts_available(), "no TH Sarabun New TTF installed")
@unittest.skipUnless(find_fixer(), "thai-font-normalize fixer not installed")
class PipelineOnSyntheticPdf(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="tpd_")
        cls.lane = Path(cls.tmp)
        (cls.lane / "sources").mkdir()
        cls.s = synth.Synth(str(cls.lane / "sources" / "synthetic.pdf"))
        cls.s.build()
        cwd = os.getcwd()
        try:
            rc = pipeline.main(["--lane", str(cls.lane), "sources/synthetic.pdf", "--name", "synth", "--title", "Synthetic"]
                               + (["--render"] if find_soffice() else []))
        finally:
            os.chdir(cwd)
        cls.rc = rc
        cls.content = json.loads((cls.lane / "work/content.json").read_text(encoding="utf-8"))
        cls.qa = json.loads((cls.lane / "work/qa.json").read_text(encoding="utf-8"))
        cls.tables = json.loads((cls.lane / "work/tables.json").read_text(encoding="utf-8"))
        cls.xl = openpyxl.load_workbook(cls.lane / "final/synth-tables.xlsx")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def paras(self, pn):
        return [b for b in self.content["pages"][str(pn)]["blocks"] if b["k"] == "p"]

    @staticmethod
    def txt(b):
        return "".join(r[0] for r in b["runs"])

    def test_pipeline_ran_clean_enough(self):
        # qa returns 1 only for page-count mismatch or Thai digits; the deliberate wrong sum is a finding, not an error
        self.assertEqual(self.rc, 0)

    def test_heading_stack_centred_with_zero_indent(self):
        heads = [b for b in self.paras(1) if self.txt(b) in ("รายงานผลการดำเนินงาน", "ประจำปีงบประมาณ 2599", "หน่วยงานตัวอย่าง")]
        self.assertEqual(len(heads), 3, [self.txt(b) for b in self.paras(1)])
        for b in heads:
            self.assertTrue(b["center"], self.txt(b))
            self.assertEqual((b["left"], b["first"]), (0, 0), self.txt(b))

    def test_thai_digits_become_arabic(self):
        allp = " ".join(self.txt(b) for b in self.paras(1))
        self.assertIn("2599", allp)
        self.assertNotRegex(allp, "[๐-๙]")

    def test_list_items_not_centred(self):
        items = [b for b in self.paras(1) if self.txt(b).split(" ")[0] in ("1.", "2.", "3.")]
        self.assertEqual(len(items), 3)
        for b in items:
            self.assertFalse(b["center"], self.txt(b))
        plain = [b for b in self.paras(1) if self.txt(b) in ("งานนโยบายและแผน", "งานสารบรรณและพัสดุ", "งานการเงินและบัญชี")]
        self.assertEqual(len(plain), 3)
        for b in plain:      # unnumbered items in a stack at one x are still list items, not a centred heading
            self.assertFalse(b["center"], self.txt(b))

    def test_decimal_item_starts_new_paragraph(self):
        ps = [self.txt(b) for b in self.paras(1)]
        self.assertTrue(any(p.startswith("3.1 ") for p in ps), ps)
        wrapped = [p for p in ps if p.startswith("การดำเนินงาน")][0]
        self.assertNotIn("3.1", wrapped)

    def test_wrap_join_follows_pdf_trailing_space(self):
        ps = [self.txt(b) for b in self.paras(1)]
        wrapped = [p for p in ps if p.startswith("การดำเนินงาน")]
        self.assertEqual(len(wrapped), 1, ps)
        self.assertEqual(wrapped[0], self.s.expected_wrap.rstrip())
        self.assertEqual(wrapped[0].count(" "), 1)  # exactly the one PDF trailing space survives (wrap 2->3)
        self.assertEqual(self.qa["wrap"]["problems"], [])
        self.assertGreater(self.qa["wrap"]["checked"], 0)

    def test_table_cells_are_numeric_in_xlsx(self):
        t = [t for t in self.tables if t["pages"] == [2]][0]
        ws = self.xl[t["id"]]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[0], "pdf_page")
        self.assertIn("จำนวนนักเรียน - ชาย", header)       # two header rows flattened "top - bottom"
        self.assertIn("จำนวนนักเรียน - หญิง", header)
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0][2], 1200)                 # '1,200' -> int
        for r in rows:
            self.assertEqual(r[0], 2)                      # pdf_page column
            for v in r[2:]:
                self.assertIsInstance(v, int)
        self.assertEqual(self.qa["numeric_as_text"], 0)

    def test_word_table_has_merged_header(self):
        blk = [b for b in self.content["pages"]["2"]["blocks"] if b["k"] == "table"][0]
        cells = [c for row in blk["rows"] for c in row]
        self.assertTrue(any(c["cs"] == 2 for c in cells), "group header should span 2 columns")
        self.assertTrue(any(c["rs"] == 2 for c in cells), "row-span header cells expected")
        doc = Document(str(self.lane / "final/synth.docx"))
        self.assertEqual(len(doc.tables), 1)

    def test_qa_flags_wrong_sum(self):
        # one wrong cell (หญิง total 828, true 827) is caught by its column check and, as a consequence, by the grand-total row check
        fails = [s for s in self.qa["sums"] if s["result"] == "FAIL"]
        col = [s for s in fails if "หญิง" in s["check"] and "stated=828" in s["detail"] and "computed=827" in s["detail"]]
        self.assertEqual(len(col), 1, fails)
        self.assertLessEqual(len(fails), 2, fails)
        self.assertFalse([s for s in fails if "ชาย" in s["check"]], fails)
        passes = [s for s in self.qa["sums"] if s["result"] == "PASS"]
        self.assertGreaterEqual(len(passes), 4)           # correct column totals and the three correct row totals pass

    def test_page_counts_and_landscape_section(self):
        self.assertEqual(self.qa["pdf_pages"], 3)
        self.assertEqual(self.qa["logical_pages"], 3)
        doc = Document(str(self.lane / "final/synth.docx"))
        self.assertEqual(len(doc.sections), 2)
        self.assertGreater(doc.sections[1].page_width, doc.sections[1].page_height)
        if find_soffice() and self.qa["render"]["pages"] is not None:
            self.assertEqual(self.qa["render"]["pages"], 3)

    def test_fonts_and_page_breaks(self):
        doc = Document(str(self.lane / "final/synth.docx"))
        xml = doc.element.xml
        self.assertIn('w:cs="TH Sarabun New"', xml)
        self.assertIn('w:bidi="th-TH"', xml)
        # page 2 starts with a pageBreakBefore paragraph (page 1 does not need one; page 3 starts a section)
        self.assertEqual(xml.count("<w:pageBreakBefore"), 1)

    def test_text_sheet_covers_all_pages(self):
        pages = {r[0] for r in self.xl["text"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(pages, {1, 2, 3})

    def test_report_written(self):
        rep = (self.lane / "work/REPORT-conversion.md").read_text(encoding="utf-8")
        self.assertIn("Sum assertions", rep)
        self.assertIn("stated=828", rep)


if __name__ == "__main__":
    unittest.main()
