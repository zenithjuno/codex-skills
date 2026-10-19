#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step 3a: work/content.json -> .docx  (1 PDF page = 1 docx page, A4, TH Sarabun New with cs/lang, Arabic digits).

Page breaks use pageBreakBefore on the first paragraph of each page; orientation changes start a new section.
Margins come from the probe layout so that PDF x positions map to the same indents in Word.
After saving, the thai-font-normalize fixer (fix-thai-font) is run on the file; failure of the fixer is an error.
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tdlib import arabic, enter_lane, find_fixer, lane_args, load_config, load_json  # noqa: E402

FONT = "TH Sarabun New"
CM = 28.35
TH = re.compile("[฀-๿]")


def rfonts(rPr):
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


def style_run(run, thai, size, bold, color=None):
    rPr = run._r.get_or_add_rPr()
    rfonts(rPr)
    if bold:
        rPr.append(OxmlElement("w:b"))
        rPr.append(OxmlElement("w:bCs"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szcs)
    lang = OxmlElement("w:lang")
    if thai:
        rPr.append(OxmlElement("w:cs"))
        lang.set(qn("w:val"), "th-TH")
        lang.set(qn("w:bidi"), "th-TH")
    else:
        lang.set(qn("w:val"), "en-US")
        lang.set(qn("w:bidi"), "th-TH")
    rPr.append(lang)
    order = ["rFonts", "b", "bCs", "i", "iCs", "color", "sz", "szCs", "cs", "lang"]
    kids = list(rPr)
    kids.sort(key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
    for k in kids:
        rPr.remove(k)
    for k in kids:
        rPr.append(k)


def segments(text):
    """split into (thai?, text) segments; spaces/punctuation stick to the preceding segment"""
    segs = []
    for ch in text:
        if ch == "\t":
            segs.append(("tab", ch))
            continue
        t = bool(TH.match(ch))
        if ch in "  " or (not ch.isalnum() and not t):
            if segs and segs[-1][0] != "tab":
                segs[-1] = (segs[-1][0], segs[-1][1] + ch)
            else:
                segs.append((False, ch))
            continue
        if segs and segs[-1][0] == t:
            segs[-1] = (t, segs[-1][1] + ch)
        else:
            segs.append((t, ch))
    return segs


def add_runs(par, runs, size, color=None):
    for text, bold in runs:
        for kind, s in segments(arabic(text)):
            if kind == "tab":
                r = par.add_run()
                r.add_tab()
                style_run(r, True, size, bold, color)
                continue
            r = par.add_run(s)
            style_run(r, kind, size, bold, color)


def fmt_par(par, size, center=False, left=0.0, first=0.0, before=0.0, tabs=None, rtab=None, line=None, after=0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line or (18 if size == 16 else round(size * 1.15 * 2) / 2))
    if center:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if left:
        pf.left_indent = Cm(left)
    if first:
        pf.first_line_indent = Cm(first)
    if rtab:
        pf.tab_stops.add_tab_stop(Cm(rtab), WD_TAB_ALIGNMENT.RIGHT)
    for t in (tabs or []):
        pf.tab_stops.add_tab_stop(Cm(t), WD_TAB_ALIGNMENT.LEFT)


def set_cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_borders(cell, val="nil"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), val)
        b.append(e)
    tcPr.append(b)


def table_props(tbl, widths_cm):
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblCellMar", "w:tblInd"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    w.set(qn("w:type"), "dxa")
    tblPr.append(w)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 10), ("left", 85), ("bottom", 10), ("right", 85)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
             "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = list(tblPr)
    kids.sort(key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
    for k in kids:
        tblPr.remove(k)
    for k in kids:
        tblPr.append(k)
    for gc, wcm in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(wcm * 567)))


def add_table(doc, blk):
    rows, widths = blk["rows"], blk["widths"]
    ncols = len(widths)
    occ, placements = {}, []
    for ri, row in enumerate(rows):
        ci = 0
        for cell in row:
            while occ.get((ri, ci)):
                ci += 1
            placements.append((ri, ci, cell))
            for dr in range(cell["rs"]):
                for dc in range(cell["cs"]):
                    occ[(ri + dr, ci + dc)] = True
            ci += cell["cs"]
    nrows = max(r + c["rs"] for r, _, c in placements)
    tbl = doc.add_table(rows=nrows, cols=ncols)
    tbl.style = doc.styles["Table Grid"]
    tbl.autofit = False
    table_props(tbl, widths)
    tsize = blk.get("size", 16)
    for ri, ci, cell in placements:
        c0 = tbl.cell(ri, ci)
        c1 = tbl.cell(ri + cell["rs"] - 1, ci + cell["cs"] - 1)
        c = c0.merge(c1) if (c0 is not c1) else c0
        c.width = Cm(sum(widths[ci:ci + cell["cs"]]))
        size = cell.get("size") or tsize
        paras = cell["p"] or [[["", False]]]
        first = True
        for pr in paras:
            par = c.paragraphs[0] if first else c.add_paragraph()
            if first:
                for r in list(par.runs):
                    r._r.getparent().remove(r._r)
            first = False
            fmt_par(par, size, center=(cell["al"] == "center"))
            if cell["al"] == "right":
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_runs(par, [(t, b) for t, b, *_ in pr], size)
        va = cell.get("va", "top")
        c.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                                "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[va]
        if cell.get("shade"):
            set_cell_shade(c, cell["shade"])
        if not cell.get("b", True):
            set_cell_borders(c, "nil")
    if blk.get("merged_project"):  # optional: remove inner horizontal borders of merged-cell blocks
        for (a, b) in blk["segs"]:
            for r in range(a, b + 1):
                for cidx in range(1, ncols):
                    tcPr = tbl.cell(r, cidx)._tc.get_or_add_tcPr()
                    bd = OxmlElement("w:tcBorders")
                    if r > a:
                        e = OxmlElement("w:top")
                        e.set(qn("w:val"), "nil")
                        bd.append(e)
                    if r < b:
                        e = OxmlElement("w:bottom")
                        e.set(qn("w:val"), "nil")
                        bd.append(e)
                    if len(bd):
                        tcPr.append(bd)
    order = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
    for tc in tbl._tbl.iter(qn("w:tc")):
        tcPr = tc.tcPr
        if tcPr is None:
            continue
        kids = list(tcPr)
        seen = {k.tag: k for k in kids}
        kids = [k for k in kids if seen[k.tag] is k]
        for k in list(tcPr):
            tcPr.remove(k)
        kids.sort(key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
        for k in kids:
            tcPr.append(k)
    return tbl


def add_footer(section):
    f = section.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for fldtype, instr in (("begin", None), (None, " PAGE "), ("end", None)):
        r = p.add_run()
        style_run(r, False, 11, False)
        if fldtype:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), fldtype)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            r._r.append(it)


def setup_section(sec, orient, layouts):
    pw, ph = 21.0, 29.7
    if orient == "L":
        pw, ph = 29.7, 21.0
        sec.orientation = WD_ORIENT.LANDSCAPE
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(pw), Cm(ph)
    lay = layouts.get(orient)
    if lay:
        left = lay["left"] / CM
        right = pw - lay["right"] / CM if orient == "P" else None
        if right is None or right < 1.0:
            right = 1.9
        sec.left_margin, sec.right_margin = Cm(round(left, 2)), Cm(round(right, 2))
    else:
        sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.5)
    sec.footer_distance = Cm(0.7)
    sec.header_distance = Cm(0.7)


def merged_layouts(probe, cfg):
    """probe layout per orientation, overridden key by key by work/overrides/config.json "layout" (probe.json is rewritten by every run)."""
    lays = {o: dict(v) for o, v in ((probe or {}).get("layout") or {}).items()}
    for o, v in ((cfg or {}).get("layout") or {}).items():
        lays.setdefault(o, {}).update(v)
    return lays


def build(content, probe, out, title, cfg=None):
    layouts = merged_layouts(probe, cfg)
    pages = content["pages"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(16)
    rfonts(st.element.get_or_add_rPr())
    rpd = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rfonts(rpd)
    for tag in ("w:sz", "w:szCs", "w:lang"):
        for e in rpd.findall(qn(tag)):
            rpd.remove(e)
    for tag in ("sz", "szCs"):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:val"), "32")
        rpd.append(e)
    lg = OxmlElement("w:lang")
    lg.set(qn("w:val"), "en-US")
    lg.set(qn("w:eastAsia"), "en-US")
    lg.set(qn("w:bidi"), "th-TH")
    rpd.append(lg)
    for s in doc.styles:
        try:
            if s.element.rPr is not None:
                rfonts(s.element.rPr)
        except Exception:
            pass
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)

    cur_orient, first_section = None, True
    n_pages = int(content["meta"]["page_count"])
    for pn in range(1, n_pages + 1):
        pg = pages[str(pn)]
        orient = pg["orient"]
        section_break = False
        if first_section:
            sec = doc.sections[0]
            setup_section(sec, orient, layouts)
            add_footer(sec)
            first_section = False
            cur_orient = orient
            section_break = True
        elif orient != cur_orient:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            setup_section(sec, orient, layouts)
            cur_orient = orient
            for p in doc.paragraphs[-1:]:
                fmt_par(p, 2, line=1)
            section_break = True
        need_break = not section_break
        blocks = pg["blocks"] or [dict(k="p", runs=[["", False]], center=False, left=0, first=0, before=0, size=16)]
        for blk in blocks:
            k = blk["k"]
            if k == "p":
                par = doc.add_paragraph()
                size = blk.get("size", 16)
                fmt_par(par, size, center=blk.get("center", False), left=blk.get("left", 0), first=blk.get("first", 0),
                        before=blk.get("before", 0), tabs=blk.get("tabs"), rtab=blk.get("rtab"))
                if need_break:
                    par.paragraph_format.page_break_before = True
                    need_break = False
                add_runs(par, [(r[0], r[1]) for r in blk["runs"]], size, blk.get("color"))
            elif k == "img":
                par = doc.add_paragraph()
                fmt_par(par, 16, center=blk.get("center", False), left=blk.get("left", 0), before=blk.get("before", 0))
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                par.paragraph_format.line_spacing = 1.0
                if need_break:
                    par.paragraph_format.page_break_before = True
                    need_break = False
                par.add_run().add_picture(blk["path"], width=Cm(blk["width_cm"]))
            elif k == "table":
                if need_break or blk.get("before"):
                    sp = doc.add_paragraph()
                    fmt_par(sp, 2, before=blk.get("before", 0) if not need_break else 0,
                            line=(blk.get("before") or 1) if not need_break else 1)
                    if need_break:
                        fmt_par(sp, 2, line=1)
                        sp.paragraph_format.page_break_before = True
                        need_break = False
                add_table(doc, blk)
    if len(doc.element.body) > 1 and doc.element.body[-2].tag == qn("w:tbl"):
        p = doc.add_paragraph()
        fmt_par(p, 2, line=1)
    doc.core_properties.title = title
    doc.core_properties.author = ""
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def run_fixer(out):
    fixer = find_fixer()
    if not fixer:
        raise SystemExit("thai-font-normalize fixer not found (~/.codex/skills/thai-font-normalize/scripts/fix-thai-font)")
    r = subprocess.run([fixer, "-i", "--no-backup", out], capture_output=True, text=True)
    print((r.stdout or "")[-300:], (r.stderr or "")[-300:])
    if r.returncode != 0:
        raise SystemExit("fix-thai-font failed on %s" % out)


def main(argv=None):
    ap = lane_args(argparse.ArgumentParser(description=__doc__))
    ap.add_argument("--out", required=True, help="output .docx path (relative to lane)")
    ap.add_argument("--title", default=None)
    a = ap.parse_args(argv)
    enter_lane(a.lane)
    cfg = load_config()
    content = load_json("work/content.json")
    probe = load_json("work/probe.json")
    title = a.title or cfg.get("title") or Path(a.out).stem
    build(content, probe, a.out, title, cfg)
    print("saved", a.out)
    run_fixer(a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
