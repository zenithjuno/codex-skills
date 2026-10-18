#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step 3b: work/tables.json + work/content.json -> .xlsx  (index / T01.. / text sheets).

Sheets: index (table_id, pdf_page, caption, sheet, rows, cols, notes) · one sheet per table
(single header row, numbers as numeric cells, pdf_page column, freeze + autofilter) ·
text (pdf_page, para_no, heading, type, text) for every paragraph in the book.
Also writes work/text_rows.json (consumed by qa.py).
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tdlib import dump_json, enter_lane, lane_args, load_json  # noqa: E402

FN = "TH Sarabun New"
F_BODY = Font(name=FN, size=14)
F_HEAD = Font(name=FN, size=14, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="5B3F9F")


def fix_pdf_page(rows, cols):
    """pdf_page cells must be ints; a multi-page string such as '44, 45' keeps its first page and the full list moves to pdf_pages_all."""
    out, extra = [], False
    for r in rows:
        r = list(r)
        v = r[0]
        if isinstance(v, str):
            pages = [int(x) for x in re.findall(r"\d+", v)]
            r[0] = pages[0]
            r.append(v if len(pages) > 1 else str(pages[0]))
            extra = True
        out.append(r)
    if extra:
        for r in out:
            if len(r) < len(cols) + 1:
                r.append(str(r[0]))
        return out, cols + ["pdf_pages_all"]
    return out, cols


def style_sheet(ws, ncols, widths, wrap_cols=()):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(ncols), max(ws.max_row, 1))
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = F_HEAD
        cell.fill = FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for cell in row:
            cell.font = F_BODY
            cell.alignment = Alignment(wrap_text=(cell.column in wrap_cols), vertical="top")


def auto_widths(cols, rows):
    widths, wrap = [], []
    for i, name in enumerate(cols):
        vals = [str(r[i]) for r in rows if i < len(r) and r[i] is not None]
        longest = max([len(name)] + [max((len(x) for x in v.split("\n")), default=0) for v in vals])
        widths.append(max(8, min(60, int(longest * 1.1) + 2)))
        if any(isinstance(r[i], str) and (len(r[i]) > 40 or "\n" in r[i]) for r in rows if i < len(r)):
            wrap.append(i + 1)
    return widths, wrap


def para_text(b):
    return "".join(r[0] for r in b["runs"]).replace("\t", " ").strip()


def is_heading(b):
    runs = [r for r in b["runs"] if r[0].strip()]
    if not runs or len(para_text(b)) > 140:
        return False
    return all(r[1] for r in runs) or (b.get("size", 16) >= 18)


def cell_text(c):
    return " ".join("".join(r[0] for r in p) for p in c["p"]).strip()


def build_text_rows(content):
    rows, heading = [], ""
    for pn in range(1, int(content["meta"]["page_count"]) + 1):
        pg = content["pages"][str(pn)]
        n = 0

        def add(typ, text):
            nonlocal n
            n += 1
            rows.append([pn, n, heading, typ, text])

        for b in pg["blocks"]:
            if b["k"] == "p":
                t = para_text(b)
                if not t:
                    continue
                if is_heading(b):
                    heading = t
                add("paragraph", t)
            elif b["k"] == "img":
                add("image", "[image]")
            elif b["k"] == "table":
                for row in b["rows"]:
                    parts = [p for p in (cell_text(c) for c in row) if p]
                    if parts:
                        add("table-row", " | ".join(parts))
    return rows


def main(argv=None):
    ap = lane_args(argparse.ArgumentParser(description=__doc__))
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    enter_lane(a.lane)
    content = load_json("work/content.json")
    tables = load_json("work/tables.json", [])
    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "index"
    ws_i.append(["table_id", "pdf_page", "caption", "sheet", "rows", "cols", "notes"])
    for t in tables:
        rows, cols = fix_pdf_page(t["rows"], t["columns"])
        ws = wb.create_sheet(t["id"])
        ws.append(cols)
        for r in rows:
            ws.append(r)
        widths, wrap = auto_widths(cols, rows)
        style_sheet(ws, len(cols), widths, wrap)
        pg = t["pages"]
        label = str(pg[0]) if len(set(pg)) == 1 else "%d-%d" % (pg[0], pg[-1])
        ws_i.append([t["id"], label, t.get("caption", ""), t["id"], len(rows), len(cols), t.get("notes", "")])
    style_sheet(ws_i, 7, [10, 10, 70, 10, 8, 8, 90], (3, 7))
    ws_t = wb.create_sheet("text")
    ws_t.append(["pdf_page", "para_no", "heading", "type", "text"])
    trs = build_text_rows(content)
    for r in trs:
        ws_t.append(r)
    style_sheet(ws_t, 5, [9, 8, 44, 11, 120], (3, 5))
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    dump_json("work/text_rows.json", trs)
    print("saved", a.out, "tables", len(tables), "text rows", len(trs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
