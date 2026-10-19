#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step 4: automatic QA.  Writes work/qa.json and work/REPORT-conversion.md.

Checks
  1  per-page non-space char counts: PDF text layer vs docx (manual/placeholder pages are only self-checks)
  2  table counts per page: PDF find_tables vs docx tables vs xlsx tables
  3  no Thai digits left in docx / xlsx; numeric xlsx cells are numbers, not text
  4  wrap-join spaces: a Thai-Thai line wrap may carry a space in the docx only if the PDF line ended with a space (and vice versa)
  5  sum assertions: rows/columns named รวม / รวมทั้งสิ้น must equal the parts (a FAIL is a discrepancy to REPORT, never to fix)
  6  figures: embedded pictures per PDF page vs pictures in the docx, plus a page-overflow height estimate (WARN)
  7  logical docx page count == PDF page count; optional LibreOffice render page count (--render)
A PASS here is evidence, not proof: it does not check indents, bold, or reading order -- see references/review-checklist.md.
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract as EX  # noqa: E402
import figlib as FL  # noqa: E402
import tablelib as TL  # noqa: E402
from tdlib import (THAI_DIGIT_RE, THAI_LETTER_RE, UNSURE_RE, arabic, dump_json, enter_lane, find_soffice,  # noqa: E402
                   lane_args, load_config, load_json, nsp)

TOL = 0.006
CM_PT = 28.35


# ------------------------------------------------------------------ docx -> logical pages
EMU_CM = 360000.0


def _pictures(el):
    """(width cm, height cm) of every inline picture inside a paragraph element."""
    out = []
    for ex in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"):
        out.append((int(ex.get("cx")) / EMU_CM, int(ex.get("cy")) / EMU_CM))
    return out


def docx_pages(path, pictures=None):
    """logical pages -> (pages, ntab).  pictures: optional list that receives, per logical page, [(w_cm, h_cm), ...]."""
    body = Document(path).element.body
    pages, ntab, first = [[]], [0], True
    pics = [[]] if pictures is not None else None

    def ptext(p):
        return "".join(t.text or "" for t in p.iter(qn("w:t")))

    for el in body:
        tag = el.tag.split("}")[1]
        if tag == "p":
            ppr = el.find(qn("w:pPr"))
            pb = ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None
            sect = ppr is not None and ppr.find(qn("w:sectPr")) is not None
            if pb and not first:
                pages.append([])
                ntab.append(0)
                if pics is not None:
                    pics.append([])
            pages[-1].append(("p", ptext(el)))
            if pics is not None:
                pics[-1] += _pictures(el)
            if sect:
                pages.append([])
                ntab.append(0)
                if pics is not None:
                    pics.append([])
            first = False
        elif tag == "tbl":
            pages[-1].append(("t", "".join(t.text or "" for t in el.iter(qn("w:t")))))
            ntab[-1] += 1
            first = False
    while len(pages) > 1 and not pages[-1]:
        pages.pop()
        ntab.pop()
        if pics is not None:
            pics.pop()
    if pics is not None:
        pictures[:] = pics
    return pages, ntab


# ------------------------------------------------------------------ figures
TOP_CM, BOTTOM_CM = 1.9, 1.5      # build_docx.setup_section margins
SPACING_CM = 6 / 28.35            # `before` of a picture paragraph


def figure_checks(doc, content, dpics, cfg):
    """per page: figures in the PDF vs pictures in the docx, and a height estimate that warns when the page would overflow.
    Estimate = PDF text/table heights (a gap a picture filled counts at most 12 pt) + docx picture heights, against the body height of the page.
    An approximation (text may reflow in Word): a WARN means look at the page, not a proven overflow."""
    c = FL.fig_cfg(cfg)
    rep = FL.repeated_signatures(doc, c["min_area"])
    rows = []
    for pn in range(1, doc.page_count + 1):
        page = doc[pn - 1]
        pg = content["pages"][str(pn)]
        figs = FL.page_figure_info(doc, page, c["min_area"], rep)
        kept, skipped = FL.select_figures(pn, figs, cfg)
        groups = FL.group_figures(kept)
        have = dpics[pn - 1] if pn - 1 < len(dpics) else []
        manual = pg["source"] in ("manual", "placeholder") or pg["cls"] == "scan"
        expected = None if manual else len(groups)
        status, note = "OK", ""
        if expected is not None and len(have) < expected:
            status, note = "CHECK", "%d picture(s) of the PDF missing in the docx" % (expected - len(have))
        elif expected is not None and len(have) > expected:
            note = "%d more picture(s) in the docx than exported figures (hook/override)" % (len(have) - expected)
        tabs = [f for f in figs if f["figure_table"]]
        if tabs:
            note = (note + "; " if note else "") + "figure-table xref %s: picture of a table, rebuild as a real table" % ",".join(str(f["xref"]) for f in tabs)
        warn = ""
        if have:
            allowed = page.rect.height / CM_PT - TOP_CM - BOTTOM_CM
            pics_h = sum(h for _, h in have) + SPACING_CM * len(have)
            other = 0.0
            if not manual:
                items = sorted([(l["y0"], l["y1"]) for l in EX.page_lines(page)] + [(t.bbox[1], t.bbox[3]) for t in EX.good_tables(page)])
                spans = [(min(f["bbox"][1] for f in g), max(f["bbox"][3] for f in g)) for g in groups]
                for k, (y0, y1) in enumerate(items):
                    other += y1 - y0
                    if k + 1 < len(items):
                        gap = max(0.0, items[k + 1][0] - y1)
                        inside = sum(max(0.0, min(b, items[k + 1][0]) - max(a, y1)) for a, b in spans)
                        other += min(max(0.0, gap - inside), 12.0) if inside else gap
                other /= CM_PT
            est = other + pics_h
            if est > allowed:
                warn = "estimated %.1f cm of content > %.1f cm body height (pictures %.1f cm): shrink a picture (figures.width_cm) or check the page in Word" % (est, allowed, pics_h)
        rows.append(dict(pdf_page=pn, pdf_figures=len(figs), skipped=len(skipped), exported=expected, docx_pictures=len(have),
                         status=status, note=note, overflow=warn))
    return rows


# ------------------------------------------------------------------ sums
def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _close(a, b):
    return abs(a - b) <= TOL + 1e-9 * max(abs(a), abs(b))


SKIP_COL = re.compile(r"^(ที่|ลำดับ|ข้อ|no\.?|#|pdf_page)$|ร้อยละ|%|เฉลี่ย|average|mean|คะแนน", re.I)


TOTAL_HEADS = ("รวม", "รวมทั้งสิ้น")
ROW_KIND_COL = "ชนิดแถว"  # optional row-type column; values starting with รวม mark total/subtotal rows


def _kind(rows, c):
    """'amt' if any value of column c has a fractional part, else 'cnt' (counts / whole numbers)."""
    for r in rows:
        v = _num(r[c]) if c < len(r) else None
        if isinstance(v, float) and v != int(v):
            return "amt"
    return "cnt"


def _head(col):
    return col.split(" - ")[-1].strip()


def _prefix(col):
    return col.rsplit(" - ", 1)[0].strip() if " - " in col else ""


def total_parts(cols, rows, numeric, c, declared=None):
    """numeric columns that make up total column c.
    declared: config qa.total_columns entries {"table", "total", "parts": [header, ...]} (matched by full or last-segment header).
    Else: columns to its left under the same group header; with no group, columns of the same kind (counts vs amounts)."""
    def match(name):
        return [i for i in numeric if name in (cols[i], _head(cols[i]))]
    for d in declared or []:
        if d["total"] in (cols[c], _head(cols[c])):
            return [i for n in d["parts"] for i in match(n) if i != c]
    left = [i for i in numeric if i < c and not _head(cols[i]).startswith("รวม")]
    pre = _prefix(cols[c])
    if pre:
        grp = [i for i in left if _prefix(cols[i]) == pre]
        if grp:
            return grp
    k = _kind(rows, c)
    return [i for i in left if _kind(rows, i) == k]


def sum_checks(tables, cfg=None):
    res = []
    declared = ((cfg or {}).get("qa") or {}).get("total_columns") or []

    def add(tid, check, ok, detail):
        res.append(dict(table=tid, check=check, result="PASS" if ok else "FAIL", detail=detail))

    for t in tables:
        cols, rows = t["columns"], t["rows"]
        if len(rows) < 2:
            continue
        tid = t["id"]

        kind_col = next((i for i, c in enumerate(cols) if c.strip() == ROW_KIND_COL), None)

        def total_label(r):
            """label of a total row. With a `ชนิดแถว` column, that column decides (a value starting with รวม);
            otherwise the FIRST text cell of the row starts with รวม (a รวม... cell elsewhere is just text)."""
            if kind_col is not None:
                k = r[kind_col] if kind_col < len(r) else None
                if not (isinstance(k, str) and k.strip().startswith("รวม")):
                    return None
                return next((v for i, v in enumerate(r) if i > 0 and i != kind_col and isinstance(v, str) and v.strip()), k)
            v = next((v for i, v in enumerate(r) if i > 0 and isinstance(v, str) and v.strip()), None)
            return v if v is not None and v.strip().startswith("รวม") else None

        def is_total_row(r):
            return total_label(r) is not None

        body = [r for r in rows if not is_total_row(r)]
        numeric = []
        for i in range(1, len(cols)):
            if SKIP_COL.search(cols[i].split(" - ")[-1].strip()) or SKIP_COL.search(cols[i].strip()):
                continue
            vals = [r[i] for r in body if i < len(r) and r[i] is not None]
            nums = [v for v in vals if _num(v) is not None]
            if not vals or len(nums) < 0.6 * len(vals):
                continue
            ints = [v for v in nums if isinstance(v, int)]
            if len(ints) == len(nums) and sorted(ints) == list(range(1, len(ints) + 1)):
                continue  # running number
            numeric.append(i)
        # total rows
        last_total = -1
        subtotals = []
        for ri, r in enumerate(rows):
            if not is_total_row(r):
                continue
            label = total_label(r)
            for c in numeric:
                v = _num(r[c]) if c < len(r) else None
                if v is None:
                    continue
                seg = [_num(x[c]) or 0 for x in rows[last_total + 1:ri] if not is_total_row(x)]
                allb = [_num(x[c]) or 0 for x in rows[:ri] if not is_total_row(x)]
                subs = [_num(rows[j][c]) or 0 for j in subtotals]
                cands = [sum(seg), sum(allb)] + ([sum(subs)] if subs else [])
                ok = any(_close(v, s) for s in cands)
                add(tid, "ผลรวมคอลัมน์ '%s' ในแถว '%s'" % (cols[c], str(label).strip()[:30]), ok,
                    "stated=%s computed=%s" % (v, " / ".join(str(round(x, 3)) for x in dict.fromkeys(cands))))
            subtotals.append(ri)
            last_total = ri
        # total columns
        for c in numeric:
            if _head(cols[c]) not in TOTAL_HEADS:
                continue
            left = total_parts(cols, rows, numeric, c, [d for d in declared if d.get("table", tid) in (tid, "*")])
            if len(left) < 2:
                continue
            for ri, r in enumerate(rows):
                v = _num(r[c]) if c < len(r) else None
                parts = [_num(r[i]) if i < len(r) else None for i in left]
                if v is None or sum(p is not None for p in parts) < 2:
                    continue
                cands = [sum(p or 0 for p in parts[k:]) for k in range(len(parts) - 1)]
                ok = any(_close(v, s) for s in cands)
                add(tid, "แถว %d: ส่วนย่อย = '%s'" % (ri + 1, cols[c]), ok, "stated=%s parts-sums=%s (p.%s)" % (v, [round(x, 3) for x in cands], r[0]))
    return res


# ------------------------------------------------------------------ wrap-join
def wrap_checks(doc, content, dpages):
    """-> (n_checked, problems).  Compares PDF line ends against the docx text of the same page."""
    problems, n = [], 0
    for pn in range(1, doc.page_count + 1):
        pg = content["pages"][str(pn)]
        if pg["source"] not in ("text-layer", "hook"):
            continue
        page = doc[pn - 1]
        ex = [(-1, t.bbox[1] - 1, 9999, t.bbox[3] + 1) for t in EX.good_tables(page)]
        lines = EX.page_lines(page, ex)
        D = arabic("\n".join(x[1] for x in dpages[pn - 1])) if pn - 1 < len(dpages) else ""
        for a, b in zip(lines, lines[1:]):
            ta = arabic("".join(r[0] for r in a["runs"]))
            tb_raw = arabic("".join(r[0] for r in b["runs"]))
            tb = tb_raw.lstrip()
            if not ta.strip() or not tb or "\t" in ta or "\t" in tb:
                continue
            ends_space = ta[-1].isspace()
            if ends_space and (ta[-2:] == "  " or tb_raw[:1].isspace()):
                continue   # a double space in the source: nothing to compare
            core = ta.rstrip()
            if not (THAI_LETTER_RE.match(core[-1]) and THAI_LETTER_RE.match(tb[0])):
                continue
            tail, head = core[-10:], tb[:10]   # long anchors: a 4-char substring also occurs elsewhere on the page (false positives)
            joined, spaced = tail + head, tail + " " + head
            n += 1
            if not ends_space and spaced in D and joined not in D:
                problems.append(dict(page=pn, kind="space-added", at="%s | %s" % (tail, head)))
            elif ends_space and joined in D and spaced not in D:
                problems.append(dict(page=pn, kind="space-lost", at="%s | %s" % (tail, head)))
    return n, problems


# ------------------------------------------------------------------ render
def render_pages(docx):
    soffice = find_soffice()
    if not soffice:
        return None, "soffice not found"
    with tempfile.TemporaryDirectory() as td:
        prof = os.path.join(td, "profile")
        try:
            subprocess.run([soffice, "--headless", "-env:UserInstallation=file://" + prof, "--convert-to", "pdf", "--outdir", td,
                            os.path.abspath(docx)], capture_output=True, timeout=300)
            pdf = os.path.join(td, Path(docx).stem + ".pdf")
            if not os.path.exists(pdf):
                return None, "LibreOffice produced no PDF"
            return fitz.open(pdf).page_count, "ok"
        except Exception as e:  # pragma: no cover
            return None, "render failed: %s" % e


# ------------------------------------------------------------------ main
def run(pdf, docx_path, xlsx_path, render=False):
    doc = fitz.open(pdf)
    content = load_json("work/content.json")
    tables = load_json("work/tables.json", [])
    dpics = []
    dpages, ntab = docx_pages(docx_path, dpics)
    xw = openpyxl.load_workbook(xlsx_path)
    xidx = {r[0]: r for r in xw["index"].iter_rows(min_row=2, values_only=True)}
    xtabs = {pn: [] for pn in range(1, doc.page_count + 1)}
    for tid, r in xidx.items():
        lab = str(r[1])
        a, b = (lab.split("-") + [lab])[:2] if "-" in lab else (lab, lab)
        for pn in range(int(a), int(b) + 1):
            if pn in xtabs:
                xtabs[pn].append(tid)
    rows = []
    for pn in range(1, doc.page_count + 1):
        pg = content["pages"][str(pn)]
        page = doc[pn - 1]
        dc = sum(nsp(x[1]) for x in dpages[pn - 1]) if pn - 1 < len(dpages) else None
        if pg["source"] in ("manual", "placeholder"):
            sc = sum(nsp("".join(r[0] for r in b["runs"])) for b in pg["blocks"] if b["k"] == "p")
            kind = "manual (self-check only)" if pg["source"] == "manual" else "TODO placeholder"
        else:
            sc = sum(nsp("".join(r[0] for r in l["runs"])) for l in EX.page_lines(page))
            kind = "text-layer" + (" + hook" if pg["source"] == "hook" else "")
        diff = (dc - sc) if dc is not None else None
        status = "OK"
        note = ""
        if pg["source"] == "placeholder":
            status, note = "TODO", "scan page not transcribed"
        elif pg["source"] == "manual":
            note = "count is the transcription itself; verify against the image"
        elif diff is None or abs(diff) > max(3, 0.005 * sc):
            status = "CHECK"
        rows.append(dict(pdf_page=pn, kind=kind, src_chars=sc, docx_chars=dc, diff=diff, pdf_tables=len(EX.good_tables(page)),
                         docx_tables=ntab[pn - 1] if pn - 1 < len(ntab) else None, xlsx_tables=",".join(xtabs[pn]) or "-",
                         status=status, note=note))
    # thai digits / numeric-as-text
    td_docx = sum(1 for pg in dpages for _, t in pg if THAI_DIGIT_RE.search(t))
    td_xlsx = sum(1 for ws in xw for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str) and THAI_DIGIT_RE.search(v))
    numtext = 0
    for ws in xw:
        if ws.title in ("index", "text"):
            continue
        names = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            numtext += sum(1 for n, v in zip(names, row) if n != "pdf_pages_all" and isinstance(v, str) and TL.parse_number(v) is not None)
    nwrap, wrap_problems = wrap_checks(doc, content, dpages)
    sums = sum_checks(tables, load_config())
    unsure = []
    for pn_, n_, h_, typ_, text in load_json("work/text_rows.json", []):
        for m in UNSURE_RE.findall(text):
            unsure.append((pn_, m))
    figures = figure_checks(doc, content, dpics, load_config())
    page_ok = len(dpages) == doc.page_count
    rendered, rnote = (render_pages(docx_path) if render else (None, "not requested"))
    out = dict(pages=rows, sums=sums, wrap=dict(checked=nwrap, problems=wrap_problems), thai_digits=dict(docx=td_docx, xlsx=td_xlsx),
               numeric_as_text=numtext, logical_pages=len(dpages), pdf_pages=doc.page_count, render=dict(pages=rendered, note=rnote),
               unsure=unsure, meta=content["meta"], figures=figures)
    out["ok"] = page_ok and td_docx == 0 and td_xlsx == 0 and numtext == 0 and not wrap_problems and all(r["status"] != "CHECK" for r in rows) \
        and all(f["status"] != "CHECK" for f in figures)
    return out


def write_report(q, path):
    A = []
    meta = q["meta"]
    fails = [s for s in q["sums"] if s["result"] == "FAIL"]
    A.append("# REPORT-conversion\n")
    A.append("Generated by thai-pdf-to-docx qa.py. PASS = evidence, not proof. Manual pages are self-checks only.\n")
    A.append("## 1. Summary\n")
    A.append("- PDF pages %d, docx logical pages %d %s" % (q["pdf_pages"], q["logical_pages"], "OK" if q["pdf_pages"] == q["logical_pages"] else "MISMATCH"))
    A.append("- Render (LibreOffice): %s (%s)" % (q["render"]["pages"] if q["render"]["pages"] is not None else "not run", q["render"]["note"]))
    A.append("- Pages: OK %d, CHECK %d, TODO %d" % tuple(sum(1 for r in q["pages"] if r["status"] == s) for s in ("OK", "CHECK", "TODO")))
    A.append("- Thai digits left: docx %d, xlsx %d; numbers stored as text in xlsx: %d" % (q["thai_digits"]["docx"], q["thai_digits"]["xlsx"], q["numeric_as_text"]))
    A.append("- Wrap-join space check: %d wraps compared, %d problems" % (q["wrap"]["checked"], len(q["wrap"]["problems"])))
    A.append("- Sum assertions: %d, FAIL %d" % (len(q["sums"]), len(fails)))
    fg = q.get("figures") or []
    A.append("- Figures: PDF %d (skipped %d), docx pictures %d; pages to look at: %d, overflow warnings: %d" % (
        sum(f["pdf_figures"] for f in fg), sum(f["skipped"] for f in fg), sum(f["docx_pictures"] for f in fg),
        sum(1 for f in fg if f["status"] != "OK" or f["note"]), sum(1 for f in fg if f["overflow"])))
    A.append("- Manual pages: %s; hook pages: %s; untranscribed scan pages: %s" % (meta["manual_pages"], meta["hook_pages"], meta["placeholder_pages"]))
    A.append("")
    A.append("## 2. Per page\n")
    A.append("| pdf page | kind | src chars | docx chars | diff | pdf tables | docx tables | xlsx | status | note |")
    A.append("|---|---|---|---|---|---|---|---|---|---|")
    for r in q["pages"]:
        A.append("| %(pdf_page)s | %(kind)s | %(src_chars)s | %(docx_chars)s | %(diff)s | %(pdf_tables)s | %(docx_tables)s | %(xlsx_tables)s | %(status)s | %(note)s |" % r)
    A.append("")
    A.append("Table counts are indicative: find_tables may count chart boxes, and tables continued across pages are one xlsx table.\n")
    A.append("## 3. Sum assertions (FAIL = the source disagrees with itself, or a cell was misread; do NOT fix, report)\n")
    for s in fails:
        A.append("- %(table)s | %(check)s | %(detail)s" % s)
    if not fails:
        A.append("- no FAIL")
    A.append("")
    A.append("## 4. Wrap-join problems\n")
    for p in q["wrap"]["problems"]:
        A.append("- p.%(page)s %(kind)s at '%(at)s'" % p)
    if not q["wrap"]["problems"]:
        A.append("- none")
    A.append("")
    A.append("## 5. Uncertain readings `[?...]`\n")
    for pn, m in q["unsure"]:
        A.append("- p.%s %s" % (pn, m))
    if not q["unsure"]:
        A.append("- none")
    A.append("")
    if meta.get("warnings"):
        A.append("## 6. Extraction warnings\n")
        for w in meta["warnings"]:
            A.append("- " + w)
        A.append("")
    A.append("## 7. Figures (pictures in the PDF vs the docx)\n")
    shown = [f for f in fg if f["pdf_figures"] or f["docx_pictures"]]
    if shown:
        A.append("| pdf page | PDF figures | skipped | exported | docx pictures | status | note / overflow estimate |")
        A.append("|---|---|---|---|---|---|---|")
        for f in shown:
            A.append("| %s | %s | %s | %s | %s | %s | %s |" % (f["pdf_page"], f["pdf_figures"], f["skipped"], "-" if f["exported"] is None else f["exported"],
                                                        f["docx_pictures"], f["status"], "; ".join(x for x in (f["note"], ("WARN " + f["overflow"]) if f["overflow"] else "") if x)))
    else:
        A.append("- no figures")
    A.append("")
    A.append("## 8. Not checked by this script\n")
    A.append("Indents/alignment against the page image, bold, reading order, scanned-page accuracy, typos kept from the source, and Word's own pagination. Use the reviewer checklist.\n")
    Path(path).write_text("\n".join(A), encoding="utf-8")


def main(argv=None):
    ap = lane_args(argparse.ArgumentParser(description=__doc__))
    ap.add_argument("pdf")
    ap.add_argument("--docx", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--render", action="store_true", help="also convert the docx with LibreOffice and compare page counts")
    a = ap.parse_args(argv)
    pdf = os.path.abspath(a.pdf)
    enter_lane(a.lane)
    q = run(pdf, a.docx, a.xlsx, a.render)
    dump_json("work/qa.json", q)
    write_report(q, "work/REPORT-conversion.md")
    print("pages OK/CHECK/TODO: %s | sum FAIL: %d | wrap problems: %d | thai digits: %s | render: %s" % (
        [sum(1 for r in q["pages"] if r["status"] == s) for s in ("OK", "CHECK", "TODO")],
        sum(1 for s in q["sums"] if s["result"] == "FAIL"), len(q["wrap"]["problems"]), q["thai_digits"], q["render"]["pages"]))
    return 0 if q["pdf_pages"] == q["logical_pages"] and not any(q["thai_digits"].values()) else 1


if __name__ == "__main__":
    sys.exit(main())
