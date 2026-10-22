#!/usr/bin/env python3
"""Smoke test: layout_pass on a tiny book + check_book catches a broken book. Run: python3 tests/test_smoke.py"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parents[1]
SCRIPTS = HERE / "scripts"


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


def main():
    tmp = Path(tempfile.mkdtemp())
    prof = json.loads((HERE / "profiles" / "pp-2570.json").read_text(encoding="utf-8"))
    (tmp / "p.json").write_text(json.dumps(prof, ensure_ascii=False), encoding="utf-8")

    # layout_pass: flat prose gets 1.25 cm + thaiDistribute, number cell centred, typo fixed
    d = Document()
    d.add_paragraph("คำนำ")
    d.add_paragraph("โรงเรียนจัดทำแผนปฏิบัติการประจำปีงบประมาณเพื่อใช้เป็นกรอบการดำเนินงานของทุกกลุ่มงานในโรงเรียน")
    d.add_paragraph("ผู้รับผิดชอบ ครูชัชพงศ์")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "รายการ", "จำนวนเงิน"
    t.cell(1, 0).text, t.cell(1, 1).text = "ค่าวัสดุ", "12,345.50"
    d.save(tmp / "in.docx")
    r = run(SCRIPTS / "layout_pass.py", tmp / "in.docx", tmp / "out.docx", "--profile", tmp / "p.json")
    assert r.returncode == 0, r.stderr
    out = Document(tmp / "out.docx")
    prose = out.paragraphs[1]._p
    assert prose.find(qn("w:pPr") + "/" + qn("w:jc")).get(qn("w:val")) == "thaiDistribute"
    assert prose.find(qn("w:pPr") + "/" + qn("w:ind")).get(qn("w:firstLine")) == "708"
    assert "ชัชพงษ์" in out.paragraphs[2].text
    cell = out.tables[0].cell(1, 1).paragraphs[0]._p
    assert cell.find(qn("w:pPr") + "/" + qn("w:jc")).get(qn("w:val")) == "center"

    # check_book: broken book must FAIL
    bad = Document()
    bad.add_paragraph("ลงชื่อ [ลายมือชื่อ]")
    run_ = bad.add_paragraph().add_run("ฉบับทำงาน")
    run_.font.color.rgb = RGBColor(0xFF, 0, 0)
    small = bad.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("99")
    small.font.size = Pt(10)
    bad.save(tmp / "bad.docx")
    pdf = fitz.open()
    for lab in ["", "", "", "1", "2", "4"]:
        pg = pdf.new_page()
        if lab:
            pg.insert_text((290, 30), lab)
        pg.insert_text((72, 200), "content line long enough to count")
    pdf.new_page()                                   # blank page
    pdf.save(tmp / "bad.pdf")
    r = run(SCRIPTS / "check_book.py", tmp / "bad.docx", tmp / "bad.pdf", "--profile", tmp / "p.json")
    assert r.returncode == 1, r.stdout
    for needle in ("FAIL หน้าว่าง", "FAIL placeholder", "FAIL ตัวแดง", "FAIL คำทำงาน", "สะดุดที่"):
        assert needle in r.stdout, (needle, r.stdout)
    print("OK", tmp)


if __name__ == "__main__":
    main()
